"""Download and parse INE (Instituto Nacional de Estadística) name statistics.

Sources (https://www.ine.es/daco/daco42/nombyapel/nombyapel.htm):
- Census: every name held by >= 20 people in Spain, per sex, with frequency and mean age.
- Births: the top-100 names given to newborns, one file per year (``nomnacYY.xlsx``).
"""

import csv
import logging
import urllib.request
from pathlib import Path

import openpyxl

from name_selector.models import Sex

log = logging.getLogger(__name__)

BASE_URL = "https://www.ine.es/daco/daco42/nombyapel"
CENSUS_FILE = "nombres_por_edad_media.xlsx"
DEFAULT_BIRTH_YEARS = (2023, 2024)
XLSX_MAGIC = b"PK\x03\x04"

CENSUS_SHEETS = {Sex.FEMALE: "Mujeres", Sex.MALE: "Hombres"}
# Columns of (name, births) in the births "TOTAL" sheet: boys on the left, girls on the right.
BIRTH_COLUMNS = {Sex.MALE: (0, 1), Sex.FEMALE: (3, 4)}


def census_csv(sex: Sex) -> str:
    return f"census_{sex.value}.csv"


def births_csv(sex: Sex) -> str:
    return f"births_{sex.value}.csv"


def births_file(year: int) -> str:
    return f"nomnac{year % 100:02d}.xlsx"


def download(url: str, dest: Path, timeout: float = 60.0) -> None:
    log.info("downloading %s", url)
    with urllib.request.urlopen(url, timeout=timeout) as response:  # noqa: S310 - fixed https host
        payload = response.read()
    if not payload.startswith(XLSX_MAGIC):
        raise ValueError(f"{url} did not return an .xlsx file (INE may have moved it)")
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(payload)


def fetch(raw_dir: Path, years: tuple[int, ...] = DEFAULT_BIRTH_YEARS) -> None:
    for name in (CENSUS_FILE, *(births_file(y) for y in years)):
        download(f"{BASE_URL}/{name}", raw_dir / name)


def parse_census(path: Path, sex: Sex = Sex.FEMALE) -> list[tuple[str, int, float]]:
    """Return ``(name, frequency, mean_age)`` for simple (single-word) names of one sex."""
    sheet = CENSUS_SHEETS[sex]
    ws = openpyxl.load_workbook(path, read_only=True)[sheet]
    rows = ws.iter_rows(values_only=True)
    for row in rows:
        if row and row[0] == "Orden":
            break
    else:
        raise ValueError(f"{path}: 'Orden' header row not found in sheet {sheet!r}")

    result = []
    for order, name, freq, age, *_ in rows:
        if not isinstance(order, int):
            break  # footnotes follow the data block
        if " " not in name:
            result.append((name.strip(), int(freq), float(age)))
    return result


def parse_births(path: Path, sex: Sex = Sex.FEMALE) -> list[tuple[str, int]]:
    """Return ``(name, births)`` for one sex from the national (``TOTAL``) sheet."""
    name_col, count_col = BIRTH_COLUMNS[sex]
    ws = openpyxl.load_workbook(path, read_only=True)["TOTAL"]
    rows = ws.iter_rows(values_only=True)
    for row in rows:
        if len(row) > name_col and row[name_col] == "TOTAL":
            break
    else:
        raise ValueError(f"{path}: {sex.value} 'TOTAL' header row not found in sheet 'TOTAL'")

    result = []
    for row in rows:
        if len(row) <= count_col:
            break
        name, count = row[name_col], row[count_col]
        if not isinstance(name, str) or not isinstance(count, int):
            break
        result.append((name.strip(), count))
    return result


def build(raw_dir: Path, processed_dir: Path, years: tuple[int, ...] = DEFAULT_BIRTH_YEARS) -> None:
    """Turn raw INE spreadsheets into the small per-sex CSVs the ranker reads."""
    processed_dir.mkdir(parents=True, exist_ok=True)
    for sex in Sex:
        census = parse_census(raw_dir / CENSUS_FILE, sex)
        with open(processed_dir / census_csv(sex), "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["name", "frequency", "mean_age"])
            writer.writerows(census)

        with open(processed_dir / births_csv(sex), "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["year", "name", "births"])
            for year in years:
                births = parse_births(raw_dir / births_file(year), sex)
                writer.writerows((year, name, n) for name, n in births)
        log.info("%s: wrote %d census names and %d birth years", sex.value, len(census), len(years))
